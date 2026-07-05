import os
import pypdf
import docx
from sqlalchemy.orm import Session
from app.db import CustomKnowledgeBase, CustomDocument, CustomDocumentChunk
from app.embeddings import get_embedding

def extract_text_from_txt(file_path: str) -> str:
    """Reads raw text from a plain text file."""
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()

def extract_text_from_pdf(file_path: str) -> str:
    """Extracts text page by page from a PDF file."""
    text_content = []
    with open(file_path, "rb") as f:
        reader = pypdf.PdfReader(f)
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_content.append(text)
    return "\n".join(text_content)

def extract_text_from_docx(file_path: str) -> str:
    """Extracts text from paragraphs and tables in a Word document."""
    doc = docx.Document(file_path)
    text_content = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_content.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_content.append(cell.text)
    return "\n".join(text_content)

def chunk_text(text: str, chunk_size: int = 1000, chunk_overlap: int = 150) -> list[str]:
    """Chunks text into overlapping segments."""
    if not text.strip():
        return []
    
    chunks = []
    start = 0
    text_len = len(text)
    
    while start < text_len:
        end = start + chunk_size
        chunk = text[start:end]
        chunks.append(chunk)
        start += chunk_size - chunk_overlap
        
    return chunks

def process_kb_document(db: Session, kb_id: str, doc_id: str) -> bool:
    """
    Extracts text, chunks, embeds, and indexes a CustomDocument.
    Updates the database status dynamically.
    """
    doc = db.query(CustomDocument).filter(CustomDocument.id == doc_id).first()
    if not doc:
        print(f"Error: CustomDocument {doc_id} not found in database.")
        return False
    
    # Update document status to processing
    doc.status = "processing"
    db.commit()
    
    try:
        file_path = doc.file_path
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found on disk at: {file_path}")
            
        ext = os.path.splitext(doc.filename)[1].lower()
        print(f"Ingesting document '{doc.filename}' (format {ext})...")
        
        # 1. Extract Text
        if ext == ".txt":
            raw_text = extract_text_from_txt(file_path)
        elif ext == ".pdf":
            raw_text = extract_text_from_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            raw_text = extract_text_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}")
            
        if not raw_text.strip():
            raise ValueError("No text content could be extracted from this document.")
            
        # 2. Chunk Text
        chunks = chunk_text(raw_text)
        print(f"Split '{doc.filename}' into {len(chunks)} chunks.")
        
        # Delete any existing chunks if re-processing
        db.query(CustomDocumentChunk).filter(CustomDocumentChunk.document_id == doc.id).delete()
        db.commit()
        
        # 3. Embed & Save Chunks
        for idx, chunk_text_block in enumerate(chunks):
            embedding = get_embedding(chunk_text_block)
            db_chunk = CustomDocumentChunk(
                document_id=doc.id,
                content=chunk_text_block,
                embedding=embedding
            )
            db.add(db_chunk)
            
        doc.status = "completed"
        doc.error_message = None
        db.commit()
        print(f"Document Ingestion Complete: '{doc.filename}'")
        
        # 4. Check KB overall status
        kb = db.query(CustomKnowledgeBase).filter(CustomKnowledgeBase.id == kb_id).first()
        if kb:
            # Check if all other docs in the KB are completed
            all_docs = db.query(CustomDocument).filter(CustomDocument.kb_id == kb_id).all()
            if all_docs and all(d.status == "completed" for d in all_docs):
                kb.status = "ready"
                db.commit()
                print(f"Vault status updated to 'ready' for KB: {kb.name}")
                
        return True
        
    except Exception as e:
        error_msg = str(e)
        print(f"Error processing custom document '{doc.filename}': {error_msg}")
        doc.status = "failed"
        doc.error_message = error_msg
        db.commit()
        
        # Update KB status to failed
        kb = db.query(CustomKnowledgeBase).filter(CustomKnowledgeBase.id == kb_id).first()
        if kb:
            kb.status = "failed"
            db.commit()
            
        return False
